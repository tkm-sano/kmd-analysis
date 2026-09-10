"""Build a source-grounded design workbook without importing an old workbook."""
from pathlib import Path
import hashlib, json, re, csv, shutil
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.pagebreak import Break

ROOT=Path(__file__).resolve().parents[1]
TARGET=ROOT/'research_model_3page_ja_revised.xlsx'
PLAN='EVRP_EXECUTION_PLAN.md'
CFG='reproducibility/config/traffic_simulation/'
DATA='reproducibility/outputs/traffic_simulation/demand/'
R12=DATA+'evrp_r12_routing/20260909_r12_routing_spec_fixture_n10_v18_geometry_reaccepted/'
R11=DATA+'evrp_r11_charging/20260909_r11_ota_primary_station_revision_model_access/'
EV=DATA+'evrp_r10_ev/20260909_r10_ev_fixture_n10_soc_full_v2/vehicle_definition.csv'
def config(n): return CFG+n
sources={}
def source(ref):
    path=ref.split('#')[0]
    p=ROOT/path
    assert p.is_file(),path
    sources[path]=hashlib.sha256(p.read_bytes()).hexdigest()
    return ref
w=Workbook();w.remove(w.active)
headers=['区分／工程','項目','定義・値・式','単位／分類','適用範囲・留意点','根拠（相対path／節）']
def sheet(name,title,subtitle):
 s=w.create_sheet(name); s.append([title]);s.merge_cells('A1:F1');s.append([subtitle]);s.merge_cells('A2:F2');s.append(headers);return s
def add(s,a,b,c,d,e,ref=PLAN):
 s.append([a,b,c,d,e,source(ref)])
 cell=s.cell(s.max_row,6);cell.hyperlink=ref.split('#')[0];cell.comment=Comment('参照ファイル SHA-256: '+sources[ref.split('#')[0]],'Source provenance')

s=sheet('01_入力と仮定','研究モデル｜入力と仮定','現行仕様の説明用資料。工程管理の正本はEVRP_EXECUTION_PLAN.md。fixture n=10は本実験の規模ではない。')
rows=[
('研究対象','大田区の住宅向けB2C配送','共通E-VRPTW条件で古典・量子最適化を比較','研究設計','需要充足率と計算資源要求を評価',PLAN+'#Objective'),
('R02–04','候補母集団 C_all','39,956地点','candidate','全件を単一EVRPとして解かない',PLAN+'#Scope'),
('R02','国勢調査2020','500m meshの世帯総数 T001141034','世帯／OBSERVED','公開統計値。meshは統計入力単位',config('evrp_r04_demand_weight_v1.yml')),
('R04','需要weight w_i','w_i = H_m / N_m','世帯相当／candidate・PROXY','配送要求q_iとは別。N_m=0は未配賦記録、隣接配賦なし',config('evrp_r04_demand_weight_v1.yml')),
('R05','customer集合 C_s','successive PPS without replacement','固定サイズn／非復元','層化なし・mesh quotaなし。nとseedは外部入力',config('evrp_r05_pps_sampling_v1.yml')),
('R05','draw probability','p_i^(k) = w_i / Σ[j∈U_k] w_j','条件付き選択確率','選択後除外。最終inclusion probabilityとは異なる',config('evrp_r05_pps_sampling_v1.yml')),
('R05–06','problem size','fixture n=10；本番nは未採択','customer','encoding・binary変数・logical qubits・資源条件確定後に本番nを採択',PLAN+'#Execution Status Snapshot'),
('R06','配送要求 q_i','全customerにq_i=1','配送要求件／ASSUMED','1 customer=1要求のモデル規約。重量・parcel countは別概念',config('evrp_r06_customer_demand_v1.yml')),
('R07','Time Window [e_i,l_i]','公開宅配受取統計から合成・較正','時間／SYNTHETIC_CALIBRATED','個別住宅の実配送ログではない。service開始に適用',config('evrp_r07_time_window_v1.yml')),
('R07','時間窓変換','指定なし／日付のみ:[0,24]；時間帯:[max(0,h−1),min(24,h+1)]；時刻proxy:[h,min(24,h+1)]','時刻原点: 当日00:00／時間','窓幅は変換規則。統計の受取時間帯そのものを許容窓としない',config('evrp_r07_time_window_v1.yml')),
('R08','service time s_i',2.5,'分/customer／ASSUMED','外部研究を参考にした仮定。travel/waiting/chargingと分離',config('evrp_r08_service_time_v1.yml')),
('R09','単一depot','DEP_006','PROXY','既存物流施設proxyを代表拠点として採択',DATA+'evrp_r09_depot/20260909_r09_depot_fixture_n10_v2/depot_definition.csv'),
('R10','車両／fleet','eCanter S；delivery；fixture 1台','台数・車種権限: ASSUMED','本番fleet sizingとは区別',config('evrp_r10_ev_v1.yml')),
('R10','payload capacity',2000,'kg／OBSERVED（PUBLIC_SPECIFICATION）','配送要求q_i=1を1kgとは解釈しない',EV),
('R10','battery capacity / catalog range','41.0 kWh / 116 km','OBSERVED（PUBLIC_SPECIFICATION）','公称仕様であり配送現場の実測ではない',EV),
('R10','energy consumption rate','41.0 / 116 ≈ 0.353448','kWh/km／COMPUTED','catalogから逆算した近似。E_ij=(d_ij[m]/1000)×rate',EV),
('R10','SOC運用','initial=1.00；minimum=0.20','比率0–1／ASSUMED','route全体・depot帰着にもminimumを適用。別のfinal閾値なし',config('evrp_r10_ev_v1.yml')),
('R10','operational_available_energy','41×(1.00−0.20)=32.8','kWh／COMPUTED','本研究の運用条件による派生量。メーカーのusable容量ではない',EV),
('R10','最大運行時間','現行fixtureでは適用無効','適用flag','設定した場合は両手法の共通制約',config('evrp_r10_ev_v1.yml')),
('R11','充電地点','京浜トラックターミナル 急速A-1/A-2','1 routing endpoint','CHAdeMO／90kW／24時間／1回30分',R11+'r11_charging_config.json'),
('R11','外部事業者の利用権','ASSUMED_AVAILABLE_FOR_MODELING','ASSUMED','設備仕様と区別。実際の入構資格を確認済みとはしない',R11+'r11_charging_config.json'),
('R11','充電受入出力','min(90,70)=70','kW／公開仕様＋モデル条件','効率1.0の仮定。30分session上限・battery上限を適用',R11+'r11_charging_config.json'),
('R12','道路network','V18 geometry-reaccepted；SUMO 1.24.0','道路d/t/a: COMPUTED','network SHA-256: 460554c7716fe5e3e1410bbee790e69745a2c423146bac88e51e3a2b95f051b2',R12+'r12_routing_config.json'),
('来歴','全入力の分類','OBSERVED / PUBLIC_STATISTICS_DERIVED / PROXY / SYNTHETIC_CALIBRATED / ASSUMED / COMPUTED','分類集合','source/hash/unit/transformation/assumption_reasonを保持',PLAN+'#Data Classification Rules')]
for row in rows:add(s,*row)

s=sheet('02_モデルと制約','研究モデル｜計算の流れと共通制約','道路path選択のtravel-time最小化と、EVRPの配送要求充足最大化を区別する。')
for row in [
('入力生成','需要生成の連鎖','公開統計 → candidate weight → PPS抽出 → q_i=1 → 時間窓 → service time','R02–R08','39,956母集団からfixture customerを選択。weightとq_iを分離',config('evrp_r05_pps_sampling_v1.yml')),
('Routing','必要OD','12 endpoint = depot 1 + customer 10 + charger 1；有向132 OD','R12','自己loop除外。candidate全件ODを生成しない',R12+'od_manifest.csv'),
('Routing','path objective','travel_time_minimizing','R12','同じ採用pathからdistanceとtravel timeを取得',R12+'r12_routing_config.json'),
('Routing','d_ij / t_ij / a_ij','road distance [m] / travel time [s] / reachable [boolean]','R12','到達不能はfalse、distance/time=null。engine failureと区別',R12+'routing_output_schema.json'),
('Routing','endpoint offset','from-node起点の有向edge上位置；出発は残余区間、到着はoffsetまで','R12','同一edge順方向は有向差。逆offsetはnetwork経由を要し、abs差で代用しない',R12+'r12_routing_config.json'),
('共通instance','I_s','customer/depot/vehicle/charger/demand/TW/service/routing/SOC/seed/hashを統合','R15設計','両手法で同一instance_id・入力・単位・constraint version',PLAN+'#Common Delivery Instance'),
('最適化','主目的','max Σ_i y_i','R17・R20設計','y_iは配送選択。全件配送を必須にしない',PLAN+'#Classical Branch'),
('最適化','第二目的','同一充足なら距離等を最小化','R17で保証方式を固定','第二目的の詳細・優先重みを未確定のまま数値化しない',PLAN+'#Classical Branch'),
]:add(s,*row)
text=(ROOT/PLAN).read_text()
block=text.split('# Common Hard Constraints\n')[1].split('\n# Data Classification Rules')[0]
for line in block.splitlines():
 if re.match(r'\| \d+ \|',line):
  a,b,c=[x.strip() for x in line.strip('|').split('|')]
  add(s,'共通制約 '+a,b,c.replace('$',''),'R16で正式凍結','両手法で共通適用・独立validatorで再計算',PLAN+'#Common Hard Constraints')
add(s,'充電式','時間とエネルギー','t_charge[min] = ΔE[kWh] / min(P_station,P_vehicle)[kW] × 60','モデル設計','t_charge≤30かつ充電後SOC≤1。時間だけを切り詰めて同量充電した扱いにしない',R11+'r11_charging_config.json')
add(s,'実装整合の留意点','R11保存式の相違','保存configには min(計算時間,30) が記載されている','記述差異・要照合','上記のsession制約式との相違を明示。本Excelでは実装修正・再受入を行わない',R11+'r11_charging_config.json')
add(s,'定義整合の留意点','q_iとpayload単位','q_iは要求件数；車両capacityはkg','未統合','件数からkgへの変換はR06で未定義。capacityとの統合を完了済みとしない',config('evrp_r06_customer_demand_v1.yml'))
add(s,'古典branch','OR-Tools','定式化 → 実行 → 独立解検証','R17–R19','feasible宣言だけでは受入しない',PLAN+'#Classical Branch')
add(s,'量子branch','QUBO → Ising → QAOA/Aer → decode','全制約encoding・penalty・離散化・変数対応を保存','R20–R24','未表現制約や非同値離散化があれば共同比較を停止',PLAN+'#Quantum Branch')
add(s,'共通検証','同一validatorで経路再計算','時間・荷量・SOC・充電・訪問・到達可能性を検査','R25','違反sampleはfeasible評価から除外',PLAN+'#Common Independent Validator')

s=sheet('03_パイプラインと評価','研究モデル｜パイプラインと評価指標','ここは設計の参照図。進行管理は計画書のみ。未実行のsolver結果・本実験DFRを数値で埋めない。')
for a,b,c in [
('R01–R04','入力監査・統計・candidate・weight','統計出典と母集団を固定し、mesh単位のweight保存則を検証'),
('R05–R08','抽出・要求・時間窓・作業時間','fixtureで入力契約・再現性を検証'),
('R09–R11','depot・EV・充電条件','位置・車両性能・SOC・利用仮定を固定'),
('R12–R14','routing仕様 → 計算 → 独立検証','endpoint/OD/path/distance/time/reachabilityを固定'),
('R15–R16','共通instance・制約','schema lockと13制約の共通validator'),
('R17–R19','OR-Tools','モデル → 実行 → 独立検証'),
('R20–R22','QUBO・等価性検証・Ising','小規模厳密fixtureとoffset込みenergy等価性'),
('R23–R25','QAOA/Aer・decode・共通検証','資源preflight後のsampleを両手法共通validatorで判定'),
('R26–R27','需要充足・古典量子比較','同一instanceで品質と計算資源を比較'),
('R28–R30','規模拡大・EV技術・最終再現性','事前予算の子runとhash/seed/command/versionの追跡')]:add(s,a,b,c,'計画','工程のPASSは本実験の完了を意味しない',PLAN+'#Pipeline Overview')
for row in [
('評価','配送要求総数','Σ_i q_i = n','件','Baseline q_i=1。fixture n=10、本番値未採択',config('evrp_r06_customer_demand_v1.yml')),
('評価','完了／未充足','完了 = Σ_i y_i；未充足 = n−Σ_i y_i','件','独立validatorに合格した解のみ。結果は未算出',PLAN+'#Classical Branch'),
('評価','配送需要充足率 DFR','Σ_i y_i / n','比率0–1','空計画がfeasibleなら0。有効解なしはN/A',PLAN+'#Valid Research Outcomes That Must NOT Stop the Whole Pipeline'),
('補助評価','走行距離・消費energy','検証済routeのdistance合計；E=d[km]×e_rate','m/kWh','道路OD値と車両route総量を区別。結果は未算出',EV),
('比較','同一条件','instance ID/customer IDs/seed/需要/TW/routing/車両条件を一致','R27','solver seedは役割別。片側だけ入力変更しない',PLAN+'#Common Delivery Instance'),
('計算資源','古典実行予算','300秒 / 8 GiB per instance','ASSUMED ceiling','実測限界値ではない',PLAN+'#Quantum Resource Preflight'),
('計算資源','量子実行予算','600秒 / 8 GiB；最大26 logical qubits','ASSUMED ceiling','binary variables / logical qubitsをencodingから求め、nで代用しない',PLAN+'#Quantum Resource Preflight'),
('計算資源','量子回路・探索設定','depth≤10,000；gates≤100,000；初期p=1、shots=1024','計画設定','回路・seed・optimizer終了理由を保存',PLAN+'#Quantum Resource Preflight'),
('規模比較','nと資源の関係','n → encoding → binary variables / logical qubits → memory/depth/gates','R20・R23・R28','本番nは未採択。25/50/100を固定scenarioにしない',PLAN+'#Scope'),
('規模比較','停止境界','last_attempted_n / largest_validated_feasible_n / first_limit_n / stop_reason','R28','停止nを最大成功nと誤記しない',PLAN+'#Problem-Size Scaling Stop Rules'),
('技術比較','EV変更対象','battery / efficiency / charging power / usable SOC','R29設計','customer・需要・窓・道路は固定。実験未実行',PLAN+'#Stage Records'),
('感度候補','service time / SOC','4–15分/customer；initial SOC 0.8–1.0、minimum 0.1–0.3','候補・未実行','正式scenarioとして固定していない',config('evrp_r10_ev_v1.yml')),
('再現性','記録対象','input/output/code/config hash、version、seed、command','R30','このExcel各根拠セルのコメントにも参照ファイルhashを保存',PLAN+'#Common Delivery Instance'),
('現在の参照','V18向けrouting','R12再固定済み。V18のR13結果は今回のExcelに掲載しない','fixture','旧V17のrouting統計をV18結果へ流用しない',PLAN+'#Execution Status Snapshot'),
('根拠の記述差異','sampling / 工程記録','計画書の旧層化記述とR05 configのPPS指定が併存','要整合確認','本Excelは利用者採択済みPPS・層化なしを記載。旧Excel項目は採用根拠にしていない',config('evrp_r05_pps_sampling_v1.yml')),
]:add(s,*row)
# Sensitivity service-time source is also explicitly retained.
s.cell(s.max_row-3,6).comment=Comment('SOC: '+sources[config('evrp_r10_ev_v1.yml')]+'\nservice time: '+source(config('evrp_r08_service_time_v1.yml'))+'\nSHA-256: '+sources[config('evrp_r08_service_time_v1.yml')],'Source provenance')
for s in w:
 s.sheet_view.showGridLines=False;s.freeze_panes='C4';s.auto_filter.ref=f'A3:F{s.max_row}'
 for col,width in zip('ABCDEF',[18,25,61,28,66,48]):s.column_dimensions[col].width=width
 s.row_dimensions[1].height=31;s.row_dimensions[2].height=36;s.row_dimensions[3].height=26
 for row in s:
  for c in row:
   c.alignment=Alignment(vertical='center',wrap_text=True);c.font=Font(name='Yu Gothic',size=10,color='243447')
   if c.row in (1,3):c.fill=PatternFill('solid',fgColor='17344D');c.font=Font(name='Yu Gothic',size=15 if c.row==1 else 10,bold=True,color='FFFFFF')
   elif c.row==2:c.fill=PatternFill('solid',fgColor='DDEBF7')
   elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F5F8')
   if c.column==6 and c.row>3:c.font=Font(name='Yu Gothic',size=9,color='146B85',underline='single')
  if row[0].row>3:s.row_dimensions[row[0].row].height=64
 s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
 s.print_options.horizontalCentered=True;s.print_area=f'A1:F{s.max_row}';s.print_title_rows='1:3'
 s.oddFooter.center.text='モデル設計の参照資料 ｜ 正本: EVRP_EXECUTION_PLAN.md';s.oddFooter.right.text='&P / &N'
w.properties.title='現行EVRPモデル設計とパイプライン';w.properties.subject='現行仕様に根拠を限定した再構成';w.properties.creator='Research model documentation'
# Backup is a byte-preserved historical workbook, never used as a content source.
backup=ROOT/'reproducibility/outputs/research_model_workbook'/datetime.now().strftime('%Y%m%d_%H%M%S')
backup.mkdir(parents=True,exist_ok=False)
if TARGET.exists():shutil.copy2(TARGET,backup/TARGET.name)
w.save(TARGET)
check=load_workbook(TARGET)
assert len(check.sheetnames)==3
for s in check:
 for row in s.iter_rows(min_row=4):
  assert all(c.value is not None for c in row)
  assert row[5].comment
texts='\n'.join(str(c.value or '') for s in check for row in s for c in row)
for forbidden in ['人件費','保守費','車両固定費','充電停止コスト','電力単価','モデル総費用','82023','82246','73547']:
 assert forbidden not in texts,forbidden
assert not any(c.data_type=='f' for s in check for row in s for c in row)
assert all(not s._charts for s in check)
report={'workbook':str(TARGET.relative_to(ROOT)),'sha256':hashlib.sha256(TARGET.read_bytes()).hexdigest(),'sources':sources,'sheets':{s.title:s.max_row-3 for s in check},'validation':'PASS: source paths exist; row provenance present; unsupported costs/legacy demand removed; no fabricated solver results, formulas or charts','prior_workbook_backup':str((backup/TARGET.name).relative_to(ROOT))}
(backup/'reconstruction_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False,indent=2))
