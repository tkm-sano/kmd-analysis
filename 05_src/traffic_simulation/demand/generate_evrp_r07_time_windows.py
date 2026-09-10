#!/usr/bin/env python3
"""Generate R07 synthetic service-start time windows from adopted receipt statistics."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import random
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
DEFAULT_CONFIG = ROOT / "reproducibility/config/traffic_simulation/evrp_r07_time_window_v1.yml"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def csv_rows(path: Path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def stats(base: Path):
    wb = load_workbook(base / "ss515_r06a.xlsx", read_only=True, data_only=True)
    ws = wb.active
    categories = {}
    for row in ws.iter_rows(min_row=6, max_row=9, values_only=True):
        categories[str(row[1])] = int(row[2])
    wb.close()
    wb = load_workbook(base / "ss508_r06a.xlsx", read_only=True, data_only=True)
    ws = wb.active
    hours = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row[1] is None or row[4] is None:
            continue
        label = str(row[3])
        if not label.endswith("時台") or not label[:-2].isdigit():
            continue
        hour = int(label.replace("時台", ""))
        hours[hour] = hours.get(hour, 0) + int(row[4])
    wb.close()
    return categories, hours


def weighted_choice(rng, items, counts):
    total = math.fsum(float(counts[x]) for x in items)
    threshold = rng.random() * total
    cumulative = 0.0
    for item in items:
        cumulative += float(counts[item])
        if threshold < cumulative:
            return item
    return items[-1]


def main() -> int:
    ap = argparse.ArgumentParser(); ap.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    ap.add_argument("--seed", type=int, required=True); ap.add_argument("--run-id", required=True); args = ap.parse_args()
    config_path = args.config.resolve(); config = yaml.safe_load(config_path.read_text(encoding="utf-8"))
    r06_path = (ROOT / config["inputs"]["r06_customer_demand"]).resolve()
    r06_manifest = (ROOT / config["inputs"]["r06_manifest"]).resolve()
    stats_manifest = (ROOT / config["inputs"]["statistics_manifest"]).resolve()
    base = (ROOT / config["inputs"]["statistics_root"]).resolve()
    output = (ROOT / config["output_root"] / args.run_id).resolve(); output.mkdir(parents=True, exist_ok=False)
    customers = csv_rows(r06_path)
    if not customers: raise ValueError("R06 customer demand is empty")
    ids = [r["stop_id"] for r in customers]
    if any(not x for x in ids) or len(ids) != len(set(ids)): raise ValueError("R06 customer IDs missing or duplicated")
    categories, hours = stats(base); cat_items = list(categories); hour_items = sorted(hours)
    rng = random.Random(args.seed); out = []
    counts = {k: 0 for k in categories}; anchor_counts = {str(h): 0 for h in hour_items}
    for order, customer in enumerate(customers, 1):
        cat_raw = weighted_choice(rng, cat_items, categories); counts[cat_raw] += 1
        if cat_raw == "指定しなかった":
            cat, e, l, width, specified = "unspecified", 0.0, 24.0, 24.0, False
        elif cat_raw == "日にち単位で指定した":
            cat, e, l, width, specified = "date_only", 0.0, 24.0, 24.0, True
        else:
            h = weighted_choice(rng, hour_items, hours); anchor_counts[str(h)] += 1
            if cat_raw == "時間帯で指定した":
                e, l, cat, specified = float(max(0, h - 1)), float(min(24, h + 1)), "time_band", True
            else:
                e, l, cat, specified = float(h), float(min(24, h + 1)), "clock_time_proxy", True
            width = l - e
        row = dict(customer)
        row.update({"time_window_category": cat, "source_time_specification_category": cat_raw,
                    "time_specified": specified, "e_i": e, "l_i": l, "window_width_hours": width,
                    "time_origin": "local_day_start_00:00", "time_unit": "hours",
                    "time_window_interpretation": "service-start allowable interval; receipt statistic is calibration input, not the interval",
                    "classification": "SYNTHETIC_CALIBRATED", "source": "MLIT Tokyo Metropolitan Goods Movement Survey 2024 ss515 and ss508",
                    "transformation": "ss515 category PPS-like categorical draw; ss508 hour weighted anchor; unspecified/date_only=[0,24], time_band=[h-1,h+1] clipped to [0,24], clock_time_proxy=[h,h+1]",
                    "assumption_reason": "Statistics do not provide customer-specific allowable windows or exact requested times; explicit proxy intervals are required for a reproducible fixture",
                    "random_seed": args.seed})
        out.append(row)
    fields = list(out[0])
    with (output / "customer_time_windows.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields); writer.writeheader(); writer.writerows(out)
    summary = {"customer_count": len(out), "category_counts": counts, "anchor_hour_counts": anchor_counts,
               "ss515_source_counts": categories, "ss508_source_hour_counts": {str(k): v for k, v in hours.items()},
               "time_origin": "local_day_start_00:00", "time_unit": "hours", "service_start_constraint": "e_i <= b_i <= l_i"}
    (output / "r07_time_window_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = {"schema_version": "evrp-r07-time-window-v1", "run_id": args.run_id, "seed": args.seed,
                "algorithm": "seeded categorical/weighted calibration draw", "rng": {"implementation": "random.Random", "version": "Python 3.11.15"},
                "classification": "SYNTHETIC_CALIBRATED", "service_start_constraint": "e_i <= b_i <= l_i",
                "inputs": {"r06_customer_demand": {"path": str(r06_path.relative_to(ROOT)), "sha256": sha256(r06_path)},
                           "r06_manifest": {"path": str(r06_manifest.relative_to(ROOT)), "sha256": sha256(r06_manifest)},
                           "statistics_manifest": {"path": str(stats_manifest.relative_to(ROOT)), "sha256": sha256(stats_manifest)},
                           "ss515": {"path": str((base / 'ss515_r06a.xlsx').relative_to(ROOT)), "sha256": sha256(base / 'ss515_r06a.xlsx')},
                           "ss508": {"path": str((base / 'ss508_r06a.xlsx').relative_to(ROOT)), "sha256": sha256(base / 'ss508_r06a.xlsx')}},
                "fixture": {"n": len(out), "is_production_problem_size": False}, "transformation": summary,
                "code": {"path": str(Path(__file__).relative_to(ROOT)), "sha256": sha256(Path(__file__))},
                "config": {"path": str(config_path.relative_to(ROOT)), "sha256": sha256(config_path)}, "outputs": {}}
    for path in sorted(output.iterdir()):
        if path.name != "r07_time_window_manifest.json": manifest["outputs"][path.name] = {"path": str(path.relative_to(ROOT)), "sha256": sha256(path)}
    (output / "r07_time_window_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"run_id": args.run_id, "customer_count": len(out), "category_counts": counts, "output": str(output.relative_to(ROOT))}, ensure_ascii=False))
    return 0


if __name__ == "__main__": raise SystemExit(main())
